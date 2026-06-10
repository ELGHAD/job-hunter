"""
Job Hunter - Excel Export Module
Creates formatted .xlsx reports from job results
"""
import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter


# ── Colors ────────────────────────────────────────────────────────────────────
HEADER_BG     = "1A1A2E"   # Dark navy
HEADER_FG     = "FFFFFF"
ACCENT        = "E94560"   # Red accent
ALT_ROW       = "F0F4FF"   # Light blue-grey
SCORE_HIGH    = "27AE60"   # Green
SCORE_MED     = "F39C12"   # Orange
SCORE_LOW     = "E74C3C"   # Red
LINK_COLOR    = "2980B9"


def _border(style="thin"):
    s = Side(style=style, color="DDDDDD")
    return Border(left=s, right=s, top=s, bottom=s)


def _score_color(score: int) -> str:
    if score >= 7:
        return SCORE_HIGH
    elif score >= 4:
        return SCORE_MED
    return SCORE_LOW


def export_to_excel(jobs: list[dict], keywords: str, output_dir: str = "exports") -> str:
    """Generate a formatted Excel file from job results. Returns file path."""
    os.makedirs(output_dir, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_kw = "".join(c if c.isalnum() else "_" for c in keywords)[:30]
    filename = f"jobs_{safe_kw}_{timestamp}.xlsx"
    filepath = os.path.join(output_dir, filename)

    wb = Workbook()

    _build_jobs_sheet(wb, jobs, keywords)
    _build_summary_sheet(wb, jobs, keywords)

    wb.save(filepath)
    return filepath


def _build_jobs_sheet(wb: Workbook, jobs: list[dict], keywords: str):
    ws = wb.active
    ws.title = "Job Listings"
    ws.freeze_panes = "A3"

    # ── Title row ─────────────────────────────────────────────────────────────
    ws.merge_cells("A1:J1")
    title_cell = ws["A1"]
    title_cell.value = f"🔍 Job Search Results — {keywords}   |   {datetime.now().strftime('%B %d, %Y')}"
    title_cell.font = Font(name="Calibri", bold=True, size=14, color=HEADER_FG)
    title_cell.fill = PatternFill("solid", fgColor=HEADER_BG)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    # ── Column headers ────────────────────────────────────────────────────────
    headers = ["#", "Job Title", "Company", "Location", "Source",
               "Work Type", "Salary", "Posted", "Score", "Apply Link"]
    col_widths = [4, 38, 24, 18, 14, 12, 18, 14, 8, 45]

    for col_idx, (header, width) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.font = Font(name="Calibri", bold=True, size=10, color=HEADER_FG)
        cell.fill = PatternFill("solid", fgColor=ACCENT)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _border()
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[2].height = 22

    # ── Data rows ─────────────────────────────────────────────────────────────
    for row_idx, job in enumerate(jobs, 3):
        is_alt = (row_idx % 2 == 0)
        row_fill = PatternFill("solid", fgColor=ALT_ROW) if is_alt else None

        row_data = [
            row_idx - 2,
            job.get("title", ""),
            job.get("company", ""),
            job.get("location", "Morocco"),
            job.get("source", ""),
            job.get("work_type", ""),
            job.get("salary", "") or "Not specified",
            job.get("posted_at", "") or "—",
            job.get("relevance_score", 0),
            job.get("url", ""),
        ]

        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = Font(name="Calibri", size=9)
            cell.alignment = Alignment(vertical="center", wrap_text=(col_idx == 2))
            cell.border = _border()
            if row_fill:
                cell.fill = row_fill

            # Score cell — color-coded
            if col_idx == 9:
                score_color = _score_color(value)
                cell.font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor=score_color)
                cell.alignment = Alignment(horizontal="center", vertical="center")

            # Link cell — hyperlink style
            if col_idx == 10 and value:
                cell.font = Font(name="Calibri", size=9, color=LINK_COLOR, underline="single")
                cell.hyperlink = value
                cell.value = "Apply →"
                cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.row_dimensions[row_idx].height = 28

    # ── Auto filter ───────────────────────────────────────────────────────────
    ws.auto_filter.ref = f"A2:J{len(jobs) + 2}"


def _build_summary_sheet(wb: Workbook, jobs: list[dict], keywords: str):
    ws = wb.create_sheet("Summary")

    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 20

    # Title
    ws.merge_cells("A1:B1")
    ws["A1"] = "📊 Search Summary"
    ws["A1"].font = Font(name="Calibri", bold=True, size=14, color=HEADER_FG)
    ws["A1"].fill = PatternFill("solid", fgColor=HEADER_BG)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    # Stats
    sources = {}
    for job in jobs:
        src = job.get("source", "Unknown")
        sources[src] = sources.get(src, 0) + 1

    stats = [
        ("Keywords searched", keywords),
        ("Total jobs found", len(jobs)),
        ("Unique sources", len(sources)),
        ("Avg relevance score", f"{sum(j.get('relevance_score', 0) for j in jobs) / max(len(jobs), 1):.1f}"),
        ("High relevance (≥7)", sum(1 for j in jobs if j.get("relevance_score", 0) >= 7)),
        ("With salary info", sum(1 for j in jobs if j.get("salary"))),
        ("Remote jobs", sum(1 for j in jobs if "remote" in str(j.get("work_type", "")).lower())),
        ("Generated at", datetime.now().strftime("%Y-%m-%d %H:%M")),
    ]

    for row_idx, (label, value) in enumerate(stats, 2):
        label_cell = ws.cell(row=row_idx, column=1, value=label)
        value_cell = ws.cell(row=row_idx, column=2, value=value)

        label_cell.font = Font(name="Calibri", bold=True, size=10)
        value_cell.font = Font(name="Calibri", size=10)
        label_cell.border = _border()
        value_cell.border = _border()

        if row_idx % 2 == 0:
            for c in [label_cell, value_cell]:
                c.fill = PatternFill("solid", fgColor=ALT_ROW)

    # Jobs by source
    row = len(stats) + 4
    ws.cell(row=row, column=1, value="Jobs by Source").font = Font(bold=True, size=10, color=HEADER_FG)
    ws.cell(row=row, column=1).fill = PatternFill("solid", fgColor=ACCENT)
    ws.cell(row=row, column=2).fill = PatternFill("solid", fgColor=ACCENT)

    for src, count in sorted(sources.items(), key=lambda x: -x[1]):
        row += 1
        ws.cell(row=row, column=1, value=src).font = Font(name="Calibri", size=9)
        ws.cell(row=row, column=2, value=count).font = Font(name="Calibri", size=9)
        ws.cell(row=row, column=1).border = _border()
        ws.cell(row=row, column=2).border = _border()